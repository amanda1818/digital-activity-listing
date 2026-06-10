"""Generate Word and PowerPoint study reports from manload results (PRD Module E).

Run from project root:   python -m reports.generate
Aggregates by role only — never an individual ranking.
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document
from docx.shared import Pt, RGBColor
from db.database import SessionLocal
from db.models import Study
from engine.manload import compute_manload

NAVY = RGBColor(0x1E, 0x27, 0x61)
NAVY_RGB = (0x1E, 0x27, 0x61)
WHITE_RGB = (0xFF, 0xFF, 0xFF)


def generate_report(session, study_id, out_path="study_report.docx"):
    study = session.get(Study, study_id)
    results = compute_manload(session, study_id, persist=False)

    doc = Document()
    h = doc.add_heading("Activity Listing — Study Report", level=0)
    h.runs[0].font.color.rgb = NAVY
    doc.add_paragraph(f"Client study: {study.name}")
    doc.add_paragraph(f"Window: {study.start_date} to {study.end_date}")
    doc.add_paragraph(
        "Method: measured activity data + in-the-moment sampling. Required time excludes "
        "non-value-added waste (rework, waiting, over-processing). Figures are reported by "
        "role; no individual is identified or ranked.")

    doc.add_heading("Manload — required vs. actual headcount", level=1)
    overall = [r for r in results if r["period"] == "overall"]
    table = doc.add_table(rows=1, cols=6)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, t in enumerate(["Role", "Required", "Actual", "Gap", "VA %", "Waste %"]):
        hdr[i].text = t
    for r in overall:
        c = table.add_row().cells
        c[0].text = r["role"]
        c[1].text = f"{r['required_headcount']:.2f}"
        c[2].text = str(r["actual_headcount"])
        c[3].text = f"{r['gap']:+.2f}"
        c[4].text = f"{r['va_pct']:.0f}"
        c[5].text = f"{r['nva_waste_pct']:.0f}"

    cyc = [r for r in results if r["period"] in ("baseline", "peak")]
    if cyc:
        doc.add_heading("Cyclical roles — baseline vs. peak load", level=1)
        doc.add_paragraph(
            "For cyclical functions, headcount is reported as a range across the cycle "
            "rather than a single average. Peak corresponds to month-end close.")
        t2 = doc.add_table(rows=1, cols=4)
        t2.style = "Light Grid Accent 1"
        for i, t in enumerate(["Role", "Period", "Required", "Actual"]):
            t2.rows[0].cells[i].text = t
        for r in sorted(cyc, key=lambda x: (x["role"], x["period"])):
            c = t2.add_row().cells
            c[0].text = r["role"]; c[1].text = r["period"]
            c[2].text = f"{r['required_headcount']:.2f}"; c[3].text = str(r["actual_headcount"])

    doc.add_heading("Recommendations", level=1)
    doc.add_paragraph("[Consultant to complete: process changes to remove the waste identified above, "
                      "and staffing decision for cyclical peaks — staff-to-peak, staff-to-average, or smooth.]")

    doc.save(out_path)
    return out_path


def generate_pptx(session, study_id, out_path="study_report.pptx"):
    """Generate a branded PowerPoint deck with live manload data (PRD §E1)."""
    from pptx import Presentation
    from pptx.util import Inches, Pt as PptPt, Emu
    from pptx.dml.color import RGBColor as PptRGB
    from pptx.enum.text import PP_ALIGN

    study = session.get(Study, study_id)
    results = compute_manload(session, study_id, persist=False)
    overall = [r for r in results if r["period"] == "overall"]
    cyc = [r for r in results if r["period"] in ("baseline", "peak")]

    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    NAVY_C = PptRGB(*NAVY_RGB)
    WHITE_C = PptRGB(*WHITE_RGB)
    GOLD_C = PptRGB(0xF4, 0xC4, 0x30)

    blank = prs.slide_layouts[6]  # completely blank

    def _add_text(slide, text, left, top, width, height, size=18, bold=False,
                  color=None, align=PP_ALIGN.LEFT, wrap=True):
        tf = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        tf.text_frame.word_wrap = wrap
        p = tf.text_frame.paragraphs[0]
        p.alignment = align
        run = p.add_run()
        run.text = text
        run.font.size = PptPt(size)
        run.font.bold = bold
        if color:
            run.font.color.rgb = color
        return tf

    def _fill_bg(slide, color_rgb):
        from pptx.util import Emu
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = PptRGB(*color_rgb)

    # Slide 1: Title
    sl = prs.slides.add_slide(blank)
    _fill_bg(sl, NAVY_RGB)
    _add_text(sl, "Activity Listing", 0.5, 1.5, 12, 1.2, size=40, bold=True,
              color=WHITE_C, align=PP_ALIGN.CENTER)
    _add_text(sl, "Manload & Efficiency Study", 0.5, 2.9, 12, 0.7, size=24,
              color=GOLD_C, align=PP_ALIGN.CENTER)
    _add_text(sl, study.name, 0.5, 3.8, 12, 0.6, size=20, color=WHITE_C, align=PP_ALIGN.CENTER)
    _add_text(sl, f"Study window: {study.start_date}  →  {study.end_date}",
              0.5, 4.5, 12, 0.5, size=14, color=WHITE_C, align=PP_ALIGN.CENTER)
    _add_text(sl, "Method: measured telemetry + in-the-moment sampling. "
              "Aggregated by role only. No individual is identified or ranked.",
              0.5, 5.5, 12, 0.8, size=11, color=WHITE_C, align=PP_ALIGN.CENTER)

    # Slide 2: Methodology summary
    sl = prs.slides.add_slide(blank)
    _fill_bg(sl, (0xF5, 0xF7, 0xFA))
    _add_text(sl, "How the data was captured", 0.5, 0.4, 12, 0.7, size=28,
              bold=True, color=NAVY_C)
    bullets = [
        "① Passive telemetry  —  calendar, email metadata, app/window activity (no content)",
        "② Active sampling (ESM)  —  3–6 short in-the-moment prompts per day (randomised)",
        "③ Idle-gap attribution  —  off-screen time categorised on return (meetings, calls, breaks)",
        "④ Classification engine  —  rules → ML → Claude API; low-confidence blocks reviewed",
        "⑤ Manload formula  —  required time (VA + necessary) ÷ available capacity per head",
    ]
    for i, b in enumerate(bullets):
        _add_text(sl, b, 0.8, 1.4 + i * 1.0, 11.5, 0.8, size=14, color=NAVY_C)

    # Slide 3: Manload summary table
    sl = prs.slides.add_slide(blank)
    _fill_bg(sl, (0xF5, 0xF7, 0xFA))
    _add_text(sl, "Required vs. Actual Headcount", 0.5, 0.3, 12, 0.7, size=28,
              bold=True, color=NAVY_C)

    col_w = [3.5, 1.5, 1.5, 1.3, 1.5, 1.5]
    headers = ["Role", "Required HC", "Actual HC", "Gap", "VA %", "Waste %"]
    tbl_top = 1.2
    row_h = 0.55
    x = 0.4
    for j, (h, w) in enumerate(zip(headers, col_w)):
        tf = sl.shapes.add_textbox(Inches(x), Inches(tbl_top), Inches(w), Inches(row_h))
        tf.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        run = tf.text_frame.paragraphs[0].add_run()
        run.text = h
        run.font.bold = True
        run.font.size = PptPt(13)
        run.font.color.rgb = WHITE_C
        from pptx.util import Pt as PPt
        fill = tf.fill; fill.solid(); fill.fore_color.rgb = NAVY_C
        x += w

    for i, r in enumerate(overall):
        row_top = tbl_top + (i + 1) * row_h
        x = 0.4
        vals = [r["role"], f"{r['required_headcount']:.2f}", str(r["actual_headcount"]),
                f"{r['gap']:+.2f}", f"{r['va_pct']:.0f}%", f"{r['nva_waste_pct']:.0f}%"]
        bg = (0xE8, 0xEB, 0xF5) if i % 2 == 0 else WHITE_RGB
        for j, (v, w) in enumerate(zip(vals, col_w)):
            tf = sl.shapes.add_textbox(Inches(x), Inches(row_top), Inches(w), Inches(row_h))
            tf.text_frame.word_wrap = False
            tf.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            run = tf.text_frame.paragraphs[0].add_run()
            run.text = v
            run.font.size = PptPt(13)
            run.font.color.rgb = NAVY_C
            fill = tf.fill; fill.solid(); fill.fore_color.rgb = PptRGB(*bg)
            x += w

    # Slide 4: VA / NVA split
    sl = prs.slides.add_slide(blank)
    _fill_bg(sl, (0xF5, 0xF7, 0xFA))
    _add_text(sl, "Time split per role  (VA / NVA-necessary / Waste)", 0.5, 0.3, 12, 0.7,
              size=24, bold=True, color=NAVY_C)
    bar_colors = [PptRGB(0x2E, 0x86, 0xAB), PptRGB(0xF4, 0xC4, 0x30), PptRGB(0xE8, 0x3A, 0x3A)]
    bar_labels = ["VA", "NVA-necessary", "Waste"]
    y_pos = 1.3
    for r in overall:
        _add_text(sl, r["role"], 0.4, y_pos, 2.0, 0.45, size=13, bold=True, color=NAVY_C)
        bar_x = 2.6
        total_w = 8.5
        for pct, col in zip([r["va_pct"], r["nva_nec_pct"], r["nva_waste_pct"]], bar_colors):
            seg_w = total_w * pct / 100
            if seg_w > 0.05:
                rect = sl.shapes.add_shape(1, Inches(bar_x), Inches(y_pos + 0.05),
                                           Inches(seg_w), Inches(0.35))
                rect.fill.solid(); rect.fill.fore_color.rgb = col
                rect.line.fill.background()
                if seg_w > 0.4:
                    tf = rect.text_frame
                    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
                    run = tf.paragraphs[0].add_run()
                    run.text = f"{pct:.0f}%"
                    run.font.size = PptPt(10)
                    run.font.color.rgb = WHITE_C
            bar_x += seg_w
        y_pos += 0.7
    lx = 2.6
    for lbl, col in zip(bar_labels, bar_colors):
        rect = sl.shapes.add_shape(1, Inches(lx), Inches(y_pos + 0.1), Inches(0.25), Inches(0.25))
        rect.fill.solid(); rect.fill.fore_color.rgb = col; rect.line.fill.background()
        _add_text(sl, lbl, lx + 0.3, y_pos + 0.05, 1.5, 0.35, size=11, color=NAVY_C)
        lx += 2.0

    # Slide 5: Cyclical roles (if any)
    if cyc:
        sl = prs.slides.add_slide(blank)
        _fill_bg(sl, (0xF5, 0xF7, 0xFA))
        _add_text(sl, "Cyclical roles — baseline vs. peak load", 0.5, 0.3, 12, 0.7,
                  size=24, bold=True, color=NAVY_C)
        _add_text(sl, "Peak = month-end close. Staffing decision: staff-to-peak (idle off-peak), "
                  "staff-to-average (overtime at peak), or smooth the peak (level-load / automate).",
                  0.5, 1.1, 12, 0.6, size=13, color=NAVY_C)
        cols2 = [3.0, 2.0, 2.0, 2.0]
        hdrs2 = ["Role", "Period", "Required HC", "Actual HC"]
        x = 0.4
        for h, w in zip(hdrs2, cols2):
            tf = sl.shapes.add_textbox(Inches(x), Inches(1.9), Inches(w), Inches(0.5))
            tf.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            run = tf.text_frame.paragraphs[0].add_run()
            run.text = h; run.font.bold = True; run.font.size = PptPt(13)
            run.font.color.rgb = WHITE_C
            fill = tf.fill; fill.solid(); fill.fore_color.rgb = NAVY_C
            x += w
        for i, r in enumerate(sorted(cyc, key=lambda x: (x["role"], x["period"]))):
            row_top = 2.5 + i * 0.5
            x = 0.4
            vals2 = [r["role"], r["period"], f"{r['required_headcount']:.2f}", str(r["actual_headcount"])]
            bg = (0xE8, 0xEB, 0xF5) if i % 2 == 0 else WHITE_RGB
            for v, w in zip(vals2, cols2):
                tf = sl.shapes.add_textbox(Inches(x), Inches(row_top), Inches(w), Inches(0.45))
                tf.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
                run = tf.text_frame.paragraphs[0].add_run()
                run.text = v; run.font.size = PptPt(13); run.font.color.rgb = NAVY_C
                fill = tf.fill; fill.solid(); fill.fore_color.rgb = PptRGB(*bg)
                x += w

    # Slide 6: Findings & recommendations
    sl = prs.slides.add_slide(blank)
    _fill_bg(sl, NAVY_RGB)
    _add_text(sl, "Key findings & recommendations", 0.5, 0.5, 12, 0.7, size=28,
              bold=True, color=WHITE_C)
    placeholders = [
        "Process changes to eliminate the highest-waste activities identified above.",
        "Staffing decision for cyclical peaks — staff-to-peak / staff-to-average / smooth.",
        "Rework root-cause analysis for activities flagged is_rework_category = true.",
        "[Consultant to complete with client-specific insights]",
    ]
    for i, txt in enumerate(placeholders):
        _add_text(sl, f"•  {txt}", 0.8, 1.6 + i * 1.0, 11.5, 0.8, size=15, color=WHITE_C)

    prs.save(out_path)
    return out_path


if __name__ == "__main__":
    s = SessionLocal()
    try:
        study = s.query(Study).order_by(Study.created_at.desc()).first()
        path = generate_report(s, study.id)
        print(f"Word report written: {path}")
        pptx_path = generate_pptx(s, study.id)
        print(f"PowerPoint report written: {pptx_path}")
    finally:
        s.close()